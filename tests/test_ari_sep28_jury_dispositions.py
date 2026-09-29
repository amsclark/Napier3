"""Iowa Legal Aid's 28 September 2026 request: the two jury wordings.

A client CRS carried two adjudication wordings neither map knew, GUILTY BY
JURY and DISMISSED BY JURY ACQUITTAL, so both fell through to OTH and alerted
as unknown. Arianna Eddy asked for "Guilty by Jury" to equal GTR and
"dismissed by jury acquittal" to equal ACQ.

GUILTY BY JURY is a trial conviction, the same as GUILTY BY COURT. DISMISSED
BY JURY ACQUITTAL starts with DISMISSED, but the jury acquitted, so it codes
ACQ and not DISM.

Every case number here is synthetic. The repository is public.
"""

import os
import sys

import pytest
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import case_parser
import crs
from test_civil_case_types import _case, _row, FULL

FELONY = '00000  FECR000000'
JUVENILE = '00000  JVJV000000'

REQUESTED = [('GUILTY BY JURY', 'GTR'),
             ('DISMISSED BY JURY ACQUITTAL', 'ACQ')]


def count(wording, statute='714.2(3)'):
    return [(statute, 'SYNTHETIC THEFT', wording, '02/02/1901')]


# -- what Iowa Legal Aid asked for -------------------------------------------

@pytest.mark.parametrize('wording, code', REQUESTED)
def test_column_g_carries_the_code_asked_for(wording, code):
    row = _row(FELONY, count(wording))
    assert row['G'] == code
    assert row['reported'] == []


@pytest.mark.parametrize('wording, code', REQUESTED)
def test_the_parser_map_agrees(wording, code):
    """Column E's suffix and column F come from case_parser, column G from
    crs. Both must give the same code for the same count."""
    assert case_parser.disposition_code(wording, FELONY) == code
    assert crs.charge_code_map[wording] == {code: crs.charge_code_map[
        'GUILTY BY COURT' if code == 'GTR' else 'ACQUITTED'][code]}


@pytest.mark.parametrize('wording, code', REQUESTED)
def test_the_dnu_prefix_is_still_stripped(wording, code):
    assert case_parser.disposition_code('DNU-' + wording) == code
    assert _row(FELONY, count('DNU-' + wording))['G'] == code


@pytest.mark.parametrize('wording, code', REQUESTED)
def test_no_note_says_napier_guessed(wording, code):
    note = _row(FELONY, count(wording))['V']
    assert wording not in note, note
    assert 'OTH' not in note, note


def test_the_real_shape_both_wordings_on_one_case():
    """The client's case carried both. The conviction speaks for the case,
    and only the convicted count's statute reaches column F."""
    sheet = load_workbook(FULL)['CASE DATA']
    case = _case(FELONY, count('GUILTY BY JURY', '714.2(3)')
                 + count('DISMISSED BY JURY ACQUITTAL', '708.2A(2)(A)'))
    unknown = crs.process_case(case, sheet, crs.FIRST_CASE_ROW)
    row = str(crs.FIRST_CASE_ROW)
    assert unknown == []
    assert sheet['G' + row].value == 'GTR'
    column_f = sheet['F' + row].value or ''
    assert '714.2(3)' in column_f, column_f
    assert '708.2A(2)(A)' not in column_f, column_f


def test_an_acquittal_alone_is_not_adjudicated():
    """ACQ is in NOT_ADJUDICATED, so the statute stays out of column F and
    off the expungement sheet as an adjudicated charge."""
    sheet = load_workbook(FULL)['CASE DATA']
    case = _case(FELONY, count('DISMISSED BY JURY ACQUITTAL', '708.2A(2)(A)'))
    crs.process_case(case, sheet, crs.FIRST_CASE_ROW)
    row = str(crs.FIRST_CASE_ROW)
    assert sheet['G' + row].value == 'ACQ'
    assert '708.2A(2)(A)' not in (sheet['F' + row].value or '')


# -- and what it must not change ---------------------------------------------

@pytest.mark.parametrize('wording, code', [('DISMISSED', 'DISM'),
                                           ('DISMISSED BY COURT', 'DISM'),
                                           ('GUILTY', 'GTR'),
                                           ('GUILTY BY COURT', 'GTR')])
def test_the_neighbouring_wordings_are_untouched(wording, code):
    assert _row(FELONY, count(wording))['G'] == code


@pytest.mark.parametrize('wording, code', REQUESTED)
def test_a_juvenile_docket_reads_them_as_it_reads_their_neighbours(
        wording, code):
    """No juvenile override was asked for. GUILTY and ACQUITTED have none,
    so these read the same on a JV docket as their neighbours do."""
    neighbour = 'GUILTY BY COURT' if code == 'GTR' else 'ACQUITTED'
    assert (_row(JUVENILE, count(wording))['G']
            == _row(JUVENILE, count(neighbour))['G'])
